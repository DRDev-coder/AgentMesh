from __future__ import annotations

import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from api.config import Settings
from api.db.models import (
    BackgroundJob,
    Document,
    DocumentChunk,
    DocumentVersion,
    KnowledgeRelease,
    Workspace,
)
from api.errors import APIError
from api.saas_schemas import (
    DocumentPreview,
    DocumentPreviewChunk,
    DocumentView,
    KnowledgeReleaseView,
)
from api.services.audit import record_audit
from api.services.object_storage import ObjectStorage
from api.services.malware import scan_document
from api.tenancy import TenantContext
from shared.schemas import SourceRecord


ALLOWED_MEDIA_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "text/markdown": ".md",
}


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    number: int | None
    text: str


def _validate_file(filename: str, media_type: str, content: bytes, limit: int) -> None:
    suffix = Path(filename).suffix.lower()
    expected_suffix = ALLOWED_MEDIA_TYPES.get(media_type)
    if expected_suffix is None or suffix != expected_suffix:
        raise APIError(
            422,
            "unsupported_document",
            "Only PDF, DOCX, TXT, and Markdown documents are supported.",
        )
    if not content:
        raise APIError(422, "empty_document", "The uploaded document is empty.")
    if len(content) > limit:
        raise APIError(413, "document_too_large", "Document exceeds the 25 MB limit.")
    if content.startswith((b"MZ", b"\x7fELF")) or b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in content:
        raise APIError(422, "unsafe_document", "The uploaded file failed security validation.")
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise APIError(422, "invalid_document", "The PDF file signature is invalid.")
    if suffix == ".docx":
        if not content.startswith(b"PK"):
            raise APIError(422, "invalid_document", "The DOCX file signature is invalid.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError
        except (zipfile.BadZipFile, ValueError) as exc:
            raise APIError(422, "invalid_document", "The DOCX file is malformed.") from exc
    if suffix in {".txt", ".md"}:
        try:
            content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise APIError(422, "invalid_document", "Text documents must use UTF-8.") from exc


def create_document(
    session: Session,
    storage: ObjectStorage,
    settings: Settings,
    context: TenantContext,
    filename: str,
    media_type: str,
    content: bytes,
) -> DocumentView:
    _validate_file(filename, media_type, content, settings.document_max_bytes)
    scan_document(settings, filename, content)
    digest = hashlib.sha256(content).hexdigest()
    duplicate = session.scalar(
        select(DocumentVersion).where(
            DocumentVersion.workspace_id == context.workspace_id,
            DocumentVersion.sha256 == digest,
        )
    )
    if duplicate is not None:
        raise APIError(409, "duplicate_document", "This document is already uploaded.")
    document = Document(
        organization_id=context.organization_id,
        workspace_id=context.workspace_id,
        name=Path(filename).stem[:255],
        status="PROCESSING",
        created_by=context.actor_id,
    )
    session.add(document)
    session.flush()
    storage_key = (
        f"organizations/{context.organization_id}/workspaces/{context.workspace_id}/"
        f"documents/{document.id}/v1/{digest}{Path(filename).suffix.lower()}"
    )
    storage.put(storage_key, content, media_type)
    version = DocumentVersion(
        organization_id=context.organization_id,
        workspace_id=context.workspace_id,
        document_id=document.id,
        version=1,
        filename=filename[:255],
        media_type=media_type,
        byte_size=len(content),
        sha256=digest,
        storage_key=storage_key,
        status="PROCESSING",
    )
    session.add(version)
    session.flush()
    session.add(
        BackgroundJob(
            organization_id=context.organization_id,
            workspace_id=context.workspace_id,
            job_type="INGEST_DOCUMENT",
            payload={"document_version_id": version.id},
        )
    )
    record_audit(
        session,
        context,
        "document.uploaded",
        "document",
        document.id,
        {"filename": filename, "version_id": version.id, "sha256": digest},
    )
    return _document_view(document, version)


def create_document_version(
    session: Session,
    storage: ObjectStorage,
    settings: Settings,
    context: TenantContext,
    document_id: str,
    filename: str,
    media_type: str,
    content: bytes,
) -> tuple[DocumentView, str]:
    _validate_file(filename, media_type, content, settings.document_max_bytes)
    scan_document(settings, filename, content)
    document = session.scalar(
        select(Document).where(
            Document.id == document_id,
            Document.organization_id == context.organization_id,
            Document.workspace_id == context.workspace_id,
            Document.status != "ARCHIVED",
        )
    )
    if document is None:
        raise APIError(404, "not_found", "Document not found.")
    digest = hashlib.sha256(content).hexdigest()
    if session.scalar(
        select(DocumentVersion.id).where(
            DocumentVersion.workspace_id == context.workspace_id,
            DocumentVersion.sha256 == digest,
        )
    ):
        raise APIError(409, "duplicate_document", "This document content already exists.")
    next_version = document.current_version + 1
    storage_key = (
        f"organizations/{context.organization_id}/workspaces/{context.workspace_id}/"
        f"documents/{document.id}/v{next_version}/{digest}{Path(filename).suffix.lower()}"
    )
    storage.put(storage_key, content, media_type)
    version = DocumentVersion(
        organization_id=context.organization_id,
        workspace_id=context.workspace_id,
        document_id=document.id,
        version=next_version,
        filename=filename[:255],
        media_type=media_type,
        byte_size=len(content),
        sha256=digest,
        storage_key=storage_key,
        status="PROCESSING",
    )
    session.add(version)
    document.current_version = next_version
    document.status = "PROCESSING"
    session.flush()
    session.add(
        BackgroundJob(
            organization_id=context.organization_id,
            workspace_id=context.workspace_id,
            job_type="INGEST_DOCUMENT",
            payload={"document_version_id": version.id},
        )
    )
    record_audit(
        session,
        context,
        "document.version_uploaded",
        "document",
        document.id,
        {"version": next_version, "version_id": version.id, "sha256": digest},
    )
    return _document_view(document, version), version.id


def _extract(version: DocumentVersion, content: bytes) -> list[ExtractedPage]:
    suffix = Path(version.filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        pages = [
            ExtractedPage(number=index + 1, text=(page.extract_text() or "").strip())
            for index, page in enumerate(reader.pages)
        ]
    elif suffix == ".docx":
        document = WordDocument(io.BytesIO(content))
        pages = [
            ExtractedPage(
                number=None,
                text="\n".join(
                    paragraph.text.strip()
                    for paragraph in document.paragraphs
                    if paragraph.text.strip()
                ),
            )
        ]
    else:
        pages = [ExtractedPage(number=None, text=content.decode("utf-8").strip())]
    if not any(page.text for page in pages):
        raise ValueError("No extractable text was found; OCR is not supported")
    return pages


def _chunk_pages(pages: list[ExtractedPage]) -> list[ExtractedPage]:
    values: list[ExtractedPage] = []
    chunk_words = 700
    overlap = 80
    for page in pages:
        words = page.text.split()
        start = 0
        while start < len(words):
            end = min(len(words), start + chunk_words)
            values.append(ExtractedPage(page.number, " ".join(words[start:end])))
            if end == len(words):
                break
            start = end - overlap
    return values


def embedding(text: str) -> list[float]:
    vector = [0.0] * 384
    tokens = re.findall(r"[a-z0-9]{2,}", text.lower())
    for token in tokens:
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        index = int.from_bytes(digest[:2], "big") % len(vector)
        sign = 1.0 if digest[2] % 2 == 0 else -1.0
        vector[index] += sign
    magnitude = math.sqrt(sum(value * value for value in vector)) or 1.0
    return [round(value / magnitude, 8) for value in vector]


def _embedding_values(value: object) -> list[float]:
    if value is None:
        return []
    if hasattr(value, "tolist"):
        value = value.tolist()
    return [float(item) for item in value]


def process_document_version(
    session: Session,
    storage: ObjectStorage,
    version_id: str,
) -> None:
    version = session.get(DocumentVersion, version_id)
    if version is None or version.status not in {"PROCESSING", "FAILED"}:
        return
    document = session.get(Document, version.document_id)
    if document is None:
        return
    job = next(
        (
            item
            for item in session.scalars(
                select(BackgroundJob).where(
                    BackgroundJob.job_type == "INGEST_DOCUMENT",
                    BackgroundJob.status.in_(["PENDING", "RETRY", "RUNNING"]),
                )
            ).all()
            if item.payload.get("document_version_id") == version.id
        ),
        None,
    )
    if job:
        job.status = "RUNNING"
        job.attempts += 1
    try:
        pages = _extract(version, storage.get(version.storage_key))
        chunks = _chunk_pages(pages)
        session.query(DocumentChunk).filter(
            DocumentChunk.document_version_id == version.id
        ).delete()
        for index, chunk in enumerate(chunks):
            session.add(
                DocumentChunk(
                    organization_id=version.organization_id,
                    workspace_id=version.workspace_id,
                    document_id=version.document_id,
                    document_version_id=version.id,
                    chunk_index=index,
                    page_number=chunk.number,
                    text=chunk.text,
                    token_count=len(chunk.text.split()),
                    embedding=embedding(chunk.text),
                    source_metadata={
                        "filename": version.filename,
                        "page": chunk.number,
                        "approved": False,
                    },
                )
            )
        version.status = "READY_FOR_REVIEW"
        version.extraction_error = None
        version.page_count = len(pages)
        document.status = "READY_FOR_REVIEW"
        if job:
            job.status = "COMPLETED"
            job.completed_at = datetime.now(timezone.utc)
            job.error = None
    except Exception as exc:
        version.status = "FAILED"
        version.extraction_error = str(exc)[:500]
        document.status = "FAILED"
        if job:
            job.status = "FAILED" if job.attempts >= 4 else "RETRY"
            job.error = str(exc)[:500]


def _document_view(
    document: Document, version: DocumentVersion | None
) -> DocumentView:
    return DocumentView(
        id=document.id,
        organization_id=document.organization_id,
        workspace_id=document.workspace_id,
        name=document.name,
        status=document.status,
        current_version=document.current_version,
        filename=version.filename if version else None,
        media_type=version.media_type if version else None,
        byte_size=version.byte_size if version else None,
        sha256=version.sha256 if version else None,
        extraction_error=version.extraction_error if version else None,
        created_at=document.created_at,
        updated_at=document.updated_at,
    )


def list_documents(
    session: Session, organization_id: str, workspace_id: str
) -> list[DocumentView]:
    documents = session.scalars(
        select(Document)
        .where(
            Document.organization_id == organization_id,
            Document.workspace_id == workspace_id,
            Document.status != "ARCHIVED",
        )
        .order_by(Document.created_at.desc())
    ).all()
    values = []
    for document in documents:
        version = session.scalar(
            select(DocumentVersion).where(
                DocumentVersion.document_id == document.id,
                DocumentVersion.version == document.current_version,
            )
        )
        values.append(_document_view(document, version))
    return values


def preview_document(
    session: Session, context: TenantContext, document_id: str
) -> DocumentPreview:
    document = session.scalar(
        select(Document).where(
            Document.id == document_id,
            Document.organization_id == context.organization_id,
            Document.workspace_id == context.workspace_id,
            Document.status != "ARCHIVED",
        )
    )
    if document is None:
        raise APIError(404, "not_found", "Document not found.")
    version = session.scalar(
        select(DocumentVersion).where(
            DocumentVersion.document_id == document.id,
            DocumentVersion.version == document.current_version,
        )
    )
    if version is None:
        raise APIError(404, "not_found", "Document version not found.")
    chunks = session.scalars(
        select(DocumentChunk)
        .where(
            DocumentChunk.organization_id == context.organization_id,
            DocumentChunk.workspace_id == context.workspace_id,
            DocumentChunk.document_version_id == version.id,
        )
        .order_by(DocumentChunk.chunk_index)
        .limit(100)
    ).all()
    return DocumentPreview(
        document_id=document.id,
        document_version_id=version.id,
        version=version.version,
        chunks=[
            DocumentPreviewChunk(
                chunk_index=chunk.chunk_index,
                page_number=chunk.page_number,
                text=chunk.text,
            )
            for chunk in chunks
        ],
    )


def publish_release(
    session: Session, context: TenantContext
) -> KnowledgeReleaseView:
    versions = session.scalars(
        select(DocumentVersion)
        .join(Document, Document.id == DocumentVersion.document_id)
        .where(
            DocumentVersion.organization_id == context.organization_id,
            DocumentVersion.workspace_id == context.workspace_id,
            DocumentVersion.status.in_(["READY_FOR_REVIEW", "PUBLISHED"]),
            Document.status != "ARCHIVED",
        )
        .order_by(DocumentVersion.document_id, DocumentVersion.version.desc())
    ).all()
    selected: dict[str, DocumentVersion] = {}
    for version in versions:
        selected.setdefault(version.document_id, version)
    if not selected:
        raise APIError(409, "knowledge_not_ready", "No processed documents are ready to publish.")
    latest = session.scalar(
        select(func.max(KnowledgeRelease.version)).where(
            KnowledgeRelease.workspace_id == context.workspace_id
        )
    ) or 0
    release = KnowledgeRelease(
        organization_id=context.organization_id,
        workspace_id=context.workspace_id,
        version=latest + 1,
        document_version_ids=[value.id for value in selected.values()],
        published_by=context.actor_id,
    )
    session.add(release)
    session.flush()
    for value in selected.values():
        value.status = "PUBLISHED"
        document = session.get(Document, value.document_id)
        if document:
            document.status = "PUBLISHED"
    workspace = session.get(Workspace, context.workspace_id)
    if workspace is None:
        raise APIError(404, "not_found", "Workspace not found.")
    workspace.active_knowledge_release_id = release.id
    record_audit(
        session,
        context,
        "knowledge.published",
        "knowledge_release",
        release.id,
        {"version": release.version, "documents": release.document_version_ids},
    )
    return KnowledgeReleaseView.model_validate(release)


def list_releases(
    session: Session, organization_id: str, workspace_id: str
) -> list[KnowledgeReleaseView]:
    values = session.scalars(
        select(KnowledgeRelease)
        .where(
            KnowledgeRelease.organization_id == organization_id,
            KnowledgeRelease.workspace_id == workspace_id,
        )
        .order_by(KnowledgeRelease.version.desc())
    ).all()
    return [KnowledgeReleaseView.model_validate(value) for value in values]


def archive_document(
    session: Session, context: TenantContext, document_id: str
) -> DocumentView:
    document = session.scalar(
        select(Document).where(
            Document.id == document_id,
            Document.organization_id == context.organization_id,
            Document.workspace_id == context.workspace_id,
        )
    )
    if document is None:
        raise APIError(404, "not_found", "Document not found.")
    versions = session.scalars(
        select(DocumentVersion).where(DocumentVersion.document_id == document.id)
    ).all()
    session.query(DocumentChunk).filter(
        DocumentChunk.organization_id == context.organization_id,
        DocumentChunk.workspace_id == context.workspace_id,
        DocumentChunk.document_id == document.id,
    ).delete(synchronize_session=False)
    for item in versions:
        session.add(
            BackgroundJob(
                organization_id=context.organization_id,
                workspace_id=context.workspace_id,
                job_type="DELETE_STORED_DOCUMENT",
                payload={"storage_key": item.storage_key},
            )
        )
    document.status = "ARCHIVED"
    record_audit(
        session,
        context,
        "document.archived",
        "document",
        document.id,
    )
    version = session.scalar(
        select(DocumentVersion).where(
            DocumentVersion.document_id == document.id,
            DocumentVersion.version == document.current_version,
        )
    )
    return _document_view(document, version)


def activate_release(
    session: Session, context: TenantContext, release_id: str
) -> KnowledgeReleaseView:
    release = session.scalar(
        select(KnowledgeRelease).where(
            KnowledgeRelease.id == release_id,
            KnowledgeRelease.organization_id == context.organization_id,
            KnowledgeRelease.workspace_id == context.workspace_id,
        )
    )
    if release is None:
        raise APIError(404, "not_found", "Knowledge release not found.")
    workspace = session.get(Workspace, context.workspace_id)
    if workspace is None:
        raise APIError(404, "not_found", "Workspace not found.")
    workspace.active_knowledge_release_id = release.id
    record_audit(
        session,
        context,
        "knowledge.release_activated",
        "knowledge_release",
        release.id,
        {"version": release.version},
    )
    return KnowledgeReleaseView.model_validate(release)


def retrieve_sources(
    session: Session,
    organization_id: str,
    workspace_id: str,
    environment: str,
    query: str,
    limit: int = 5,
) -> tuple[list[SourceRecord], str | None]:
    workspace = session.get(Workspace, workspace_id)
    if workspace is None or workspace.organization_id != organization_id:
        return [], None
    release_id = workspace.active_knowledge_release_id
    version_ids: list[str] = []
    if environment == "live":
        release = session.get(KnowledgeRelease, release_id) if release_id else None
        version_ids = list(release.document_version_ids) if release else []
    else:
        version_ids = list(
            session.scalars(
                select(DocumentVersion.id).where(
                    DocumentVersion.organization_id == organization_id,
                    DocumentVersion.workspace_id == workspace_id,
                    DocumentVersion.status.in_(["READY_FOR_REVIEW", "PUBLISHED"]),
                )
            ).all()
        )
    if not version_ids:
        return [], release_id
    chunks = session.scalars(
        select(DocumentChunk).where(
            DocumentChunk.organization_id == organization_id,
            DocumentChunk.workspace_id == workspace_id,
            DocumentChunk.document_version_id.in_(version_ids),
        )
    ).all()
    query_tokens = set(re.findall(r"[a-z0-9]{2,}", query.lower()))
    query_embedding = embedding(query)
    ranked: list[tuple[float, DocumentChunk]] = []
    for chunk in chunks:
        chunk_tokens = set(re.findall(r"[a-z0-9]{2,}", chunk.text.lower()))
        lexical = len(query_tokens & chunk_tokens) / max(len(query_tokens), 1)
        vector = _embedding_values(chunk.embedding)
        cosine = sum(a * b for a, b in zip(query_embedding, vector, strict=False))
        score = max(0.0, min(1.0, (0.65 * lexical) + (0.35 * max(0.0, cosine))))
        if score > 0.05:
            ranked.append((score, chunk))
    ranked.sort(key=lambda item: (-item[0], item[1].id))
    sources = []
    for score, chunk in ranked[:limit]:
        metadata = dict(chunk.source_metadata)
        metadata["approved"] = True
        sources.append(
            SourceRecord(
                document_id=chunk.document_id,
                chunk_id=chunk.id,
                text=chunk.text,
                metadata=metadata,
                distance_or_similarity=round(score, 4),
            )
        )
    return sources, release_id
